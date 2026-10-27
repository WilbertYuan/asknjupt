from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup, NavigableString, Tag

from . import config
from .db import add_quality, init_db, insert_retrieval_unit
from .util import extract_academic_year, normalize_ws, parse_publish_date, primary_year

NOISE_SELECTORS = [
    "script",
    "style",
    "noscript",
    "nav",
    "footer",
    "header",
    ".nav",
    ".footer",
    ".header",
    "#footer",
    "#header",
    ".sidebar",
    ".right",
    ".copyright",
]


def count_tokens(text: str) -> int:
    """Prefer Qwen3 tokenizer; fall back to conservative estimate."""
    tok_json = config.TOKENIZER_DIR / "tokenizer.json"
    if not hasattr(count_tokens, "_enc"):
        count_tokens._enc = None
        if tok_json.exists():
            try:
                from tokenizers import Tokenizer

                count_tokens._enc = Tokenizer.from_file(str(tok_json))
            except Exception:
                count_tokens._enc = None
    enc = count_tokens._enc
    if enc is not None:
        try:
            return len(enc.encode(text).ids)
        except Exception:
            pass
    # Qwen-like: CJK often 1 token/char, mixed text a bit less
    cjk = sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")
    other = max(0, len(text) - cjk)
    return cjk + max(1, other // 3)


def split_chunks(text: str, max_tokens: int | None = None, overlap: int | None = None) -> list[str]:
    max_tokens = max_tokens or config.PARSE_CHUNK_TOKENS
    overlap = overlap or config.PARSE_CHUNK_OVERLAP
    paras = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    chunks: list[str] = []
    buf: list[str] = []
    buf_tok = 0
    for p in paras:
        t = count_tokens(p)
        if buf and buf_tok + t > max_tokens:
            chunks.append("\n".join(buf))
            # overlap last paragraph
            if overlap and buf:
                last = buf[-1]
                buf = [last, p] if count_tokens(last) < overlap * 3 else [p]
                buf_tok = sum(count_tokens(x) for x in buf)
            else:
                buf = [p]
                buf_tok = t
        else:
            buf.append(p)
            buf_tok += t
    if buf:
        chunks.append("\n".join(buf))
    return chunks or ([text] if text.strip() else [])


def cell_text(cell: Tag) -> str:
    return normalize_ws(cell.get_text(" ", strip=True))


def expand_html_table(table: Tag) -> dict[str, Any]:
    """Expand rowspan/colspan into a dense matrix."""
    grid: list[list[str | None]] = []
    occupy: dict[tuple[int, int], str] = {}
    r = 0
    for tr in table.find_all("tr"):
        cells = tr.find_all(["td", "th"], recursive=False)
        if not cells:
            cells = tr.find_all(["td", "th"])
        # skip occupied columns
        c = 0
        while True:
            while (r, c) in occupy:
                c += 1
            if c >= 0:
                break
        for cell in cells:
            while (r, c) in occupy:
                c += 1
            rs = int(cell.get("rowspan") or 1)
            cs = int(cell.get("colspan") or 1)
            val = cell_text(cell)
            for i in range(rs):
                for j in range(cs):
                    occupy[(r + i, c + j)] = val
            c += cs
        r += 1
    if not occupy:
        return {"headers": [], "rows": [], "raw": [], "n_rows": 0, "n_cols": 0}
    max_r = max(i for i, _ in occupy) + 1
    max_c = max(j for _, j in occupy) + 1
    matrix = [[occupy.get((i, j), "") or "" for j in range(max_c)] for i in range(max_r)]
    header_idx = 0
    best = -1
    keys = ("首选", "再选", "专业", "计划", "面积", "序号", "名称", "学费", "选科")
    for i, row in enumerate(matrix):
        joined = "".join(row)
        distinct = len({c for c in row if c})
        # skip full-row title (all cells identical)
        if distinct <= 1 and len(row) > 1:
            continue
        score = sum(1 for k in keys if k in joined) + distinct * 0.1
        if score > best:
            best = score
            header_idx = i
    headers = matrix[header_idx] if matrix else []
    body = matrix[header_idx + 1 :] if len(matrix) > header_idx + 1 else []
    return {
        "headers": headers,
        "rows": body,
        "raw": matrix,
        "n_rows": len(matrix),
        "n_cols": max_c,
        "header_row": header_idx,
    }


def extract_html(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    soup = BeautifulSoup(raw, "lxml")
    title = ""
    if soup.title and soup.title.string:
        title = normalize_ws(soup.title.get_text())
    h1 = soup.find("h1")
    if h1:
        title = normalize_ws(h1.get_text()) or title
    published = None
    meta = soup.find(class_=re.compile(r"(arti_metas|article-meta|info|pubtime|time)"))
    if meta:
        published = parse_publish_date(meta.get_text())
    if not published:
        published = parse_publish_date(soup.get_text(" ", strip=True)[:1500])

    footer_text = ""
    for tag in soup.find_all(["footer", "div"], class_=re.compile(r"footer|copyright|address", re.I)):
        footer_text += " " + normalize_ws(tag.get_text(" ", strip=True))
    body_all = soup.get_text(" ", strip=True)
    for m in re.finditer(r"(仙林校区地址[^。]{0,40}文苑路9号|锁金村校区地址[^。]{0,40}龙蟠路177号|三牌楼校区地址[^。]{0,40}新模范马路66号|咨询电话[:：\s]*[（(]?025[)）]?\s*8586666\d)", body_all):
        footer_text += " " + m.group(0)
    for sel in NOISE_SELECTORS:
        for tag in soup.select(sel):
            tag.decompose()

    content = (
        soup.select_one("div.wp_articlecontent")
        or soup.select_one("div.article")
        or soup.select_one("div#content")
        or soup.select_one("div.content")
        or soup.select_one("td.content")
        or soup.body
        or soup
    )

    tables = []
    for i, tb in enumerate(content.find_all("table")):
        expanded = expand_html_table(tb)
        caption = ""
        cap = tb.find("caption")
        if cap:
            caption = cell_text(cap)
        prev = tb.find_previous(["p", "h2", "h3", "h4", "strong"])
        if prev:
            caption = caption or normalize_ws(prev.get_text())[:80]
        tables.append(
            {
                "index": i,
                "locator": f"html-table-{i}",
                "title": caption,
                **expanded,
            }
        )
        tb.decompose()

    # headings + paragraphs
    blocks: list[dict[str, Any]] = []
    heading = title
    buf: list[str] = []
    ordinal = 0

    def flush() -> None:
        nonlocal buf, ordinal
        text = "\n".join(x for x in buf if x)
        if text.strip():
            blocks.append(
                {
                    "ordinal": ordinal,
                    "block_type": "section",
                    "heading": heading,
                    "text": text,
                    "locator": f"html-block-{ordinal}",
                }
            )
            ordinal += 1
        buf = []

    for el in content.find_all(["h1", "h2", "h3", "h4", "p", "li", "div"], recursive=True):
        if el.name in {"h1", "h2", "h3", "h4"}:
            flush()
            heading = normalize_ws(el.get_text())
            continue
        if el.name == "div" and el.find(["p", "div", "table", "h2", "h3"]):
            continue
        t = normalize_ws(el.get_text(" ", strip=True))
        if t and t not in buf:
            buf.append(t)
    flush()
    footer_text = normalize_ws(footer_text)
    if footer_text:
        blocks.append(
            {
                "ordinal": ordinal,
                "block_type": "footer",
                "heading": "页脚联系信息",
                "text": footer_text,
                "locator": "html-footer",
            }
        )

    full_text = "\n\n".join(f"{b['heading']}\n{b['text']}" if b["heading"] else b["text"] for b in blocks)
    return {
        "title": title,
        "published_at": published,
        "text": full_text,
        "blocks": blocks,
        "tables": tables,
        "doc_type": "html",
    }


def extract_image(path: Path, content_hash: str = "") -> dict[str, Any]:
    from .ocr import archive_ocr, ocr_available, ocr_image_file

    if not ocr_available():
        return {
            "title": path.stem,
            "published_at": None,
            "text": "",
            "blocks": [],
            "tables": [],
            "doc_type": "image",
            "unparsed": True,
            "ocr": False,
        }
    rec = ocr_image_file(path)
    if content_hash:
        archive_ocr(content_hash, [rec], str(path))
    text = rec.get("text") or ""
    return {
        "title": path.stem,
        "published_at": parse_publish_date(text[:800]),
        "text": text,
        "blocks": [
            {
                "ordinal": 0,
                "block_type": "ocr_image",
                "heading": path.stem,
                "text": text,
                "locator": "ocr-image",
            }
        ],
        "tables": [],
        "doc_type": "image",
        "ocr": True,
        "ocr_conf": rec.get("mean_conf"),
        "unparsed": not bool(text.strip()),
    }


def extract_pdf(path: Path) -> dict[str, Any]:
    text_pages: list[str] = []
    tables: list[dict[str, Any]] = []
    title = path.stem
    # pypdf first: some official reports crash PyMuPDF
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        meta_title = getattr(reader.metadata, "title", None) if reader.metadata else None
        if meta_title:
            title = str(meta_title) or title
        for page in reader.pages:
            text_pages.append(page.extract_text() or "")
    except Exception:
        text_pages = []
    if sum(len(t) for t in text_pages) < 80:
        try:
            import fitz

            doc = fitz.open(path)
            if doc.metadata and doc.metadata.get("title"):
                title = doc.metadata["title"] or title
            text_pages = []
            for i, page in enumerate(doc):
                text_pages.append(page.get_text("text") or "")
                try:
                    found = page.find_tables()
                    tabs = found.tables if found else []
                except Exception:
                    tabs = []
                for ti, tab in enumerate(tabs):
                    try:
                        raw = tab.extract()
                    except Exception:
                        continue
                    matrix = [[normalize_ws(str(c or "")) for c in row] for row in raw]
                    if not matrix:
                        continue
                    tables.append(
                        {
                            "index": len(tables),
                            "locator": f"pdf-p{i+1}-t{ti}",
                            "title": "",
                            "headers": matrix[0],
                            "rows": matrix[1:],
                            "raw": matrix,
                            "n_rows": len(matrix),
                            "n_cols": max(len(r) for r in matrix),
                            "pdf_page": i + 1,
                        }
                    )
            doc.close()
        except Exception:
            pass

    full = "\n\n".join(text_pages)
    used_ocr = False
    try:
        from .ocr import archive_ocr, ocr_available, ocr_pdf_pages, pdf_needs_ocr

        n_pages = max(1, len(text_pages))
        if ocr_available() and pdf_needs_ocr(full, n_pages):
            ocr_pages = ocr_pdf_pages(path, max_pages=min(max(n_pages, 1), 120))
            used_ocr = True
            digest = path.stem if len(path.stem) >= 32 else ""
            if digest:
                archive_ocr(digest, ocr_pages, str(path))
            text_pages = [p.get("text") or native for p, native in zip(ocr_pages, text_pages)] or text_pages
            full = "\n\n".join(text_pages)
    except Exception:
        used_ocr = False
    blocks = []
    for i, page_text in enumerate(text_pages):
        if normalize_ws(page_text):
            blocks.append(
                {
                    "ordinal": i,
                    "block_type": "ocr_pdf_page" if used_ocr else "pdf_page",
                    "heading": f"第{i+1}页",
                    "text": page_text,
                    "locator": f"{'ocr-p' if used_ocr else 'pdf-page-'}{i+1}",
                    "pdf_page": i + 1,
                }
            )
    return {
        "title": title,
        "published_at": parse_publish_date(full[:2000]),
        "text": full,
        "blocks": blocks,
        "tables": tables,
        "doc_type": "pdf",
        "ocr": used_ocr,
    }


def extract_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    doc = Document(str(path))
    paras = [normalize_ws(p.text) for p in doc.paragraphs if normalize_ws(p.text)]
    tables = []
    for i, tb in enumerate(doc.tables):
        matrix = [[normalize_ws(cell.text) for cell in row.cells] for row in tb.rows]
        tables.append(
            {
                "index": i,
                "locator": f"docx-table-{i}",
                "title": "",
                "headers": matrix[0] if matrix else [],
                "rows": matrix[1:] if len(matrix) > 1 else [],
                "raw": matrix,
                "n_rows": len(matrix),
                "n_cols": max((len(r) for r in matrix), default=0),
            }
        )
    text = "\n".join(paras)
    return {
        "title": path.stem,
        "published_at": parse_publish_date(text[:1500]),
        "text": text,
        "blocks": [
            {
                "ordinal": 0,
                "block_type": "docx",
                "heading": path.stem,
                "text": text,
                "locator": "docx-body",
            }
        ],
        "tables": tables,
        "doc_type": "docx",
    }


def extract_doc(path: Path) -> dict[str, Any]:
    import subprocess
    import tempfile

    text = ""
    try:
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "out.txt"
            subprocess.run(
                ["textutil", "-convert", "txt", "-output", str(out), str(path)],
                check=True,
                capture_output=True,
                timeout=60,
            )
            text = out.read_text(encoding="utf-8", errors="replace")
    except Exception:
        try:
            proc = subprocess.run(
                ["antiword", str(path)],
                capture_output=True,
                text=True,
                timeout=60,
            )
            if proc.returncode == 0:
                text = proc.stdout or ""
        except Exception:
            text = ""
    text = normalize_ws(text) if text and "\n" not in text else text
    return {
        "title": path.stem,
        "published_at": parse_publish_date((text or "")[:1500]),
        "text": text or "",
        "blocks": [
            {
                "ordinal": 0,
                "block_type": "doc",
                "heading": path.stem,
                "text": text or "",
                "locator": "doc-body",
            }
        ]
        if text
        else [],
        "tables": [],
        "doc_type": "doc",
        "unparsed": not bool(text),
    }


def extract_xls(path: Path) -> dict[str, Any]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    tables = []
    texts = []
    for si in range(book.nsheets):
        ws = book.sheet_by_index(si)
        matrix = []
        for ri in range(ws.nrows):
            matrix.append([normalize_ws("" if c is None else str(c)) for c in ws.row_values(ri)])
        matrix = [r for r in matrix if any(r)]
        tables.append(
            {
                "index": si,
                "locator": f"xls-{ws.name}",
                "title": ws.name,
                "headers": matrix[0] if matrix else [],
                "rows": matrix[1:] if len(matrix) > 1 else [],
                "raw": matrix,
                "n_rows": len(matrix),
                "n_cols": max((len(r) for r in matrix), default=0),
            }
        )
        texts.append(ws.name + "\n" + "\n".join("\t".join(r) for r in matrix[:80]))
    text = "\n\n".join(texts)
    return {
        "title": path.stem,
        "published_at": None,
        "text": text,
        "blocks": [{"ordinal": 0, "block_type": "xls", "heading": path.stem, "text": text, "locator": "xls-body"}] if text else [],
        "tables": tables,
        "doc_type": "xls",
        "unparsed": not bool(text.strip()),
    }


def extract_xlsx(path: Path) -> dict[str, Any]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    tables = []
    texts = []
    for si, name in enumerate(wb.sheetnames):
        ws = wb[name]
        matrix = []
        for row in ws.iter_rows(values_only=True):
            matrix.append([normalize_ws("" if c is None else str(c)) for c in row])
        # trim empty
        matrix = [r for r in matrix if any(r)]
        tables.append(
            {
                "index": si,
                "locator": f"xlsx-{name}",
                "title": name,
                "headers": matrix[0] if matrix else [],
                "rows": matrix[1:] if len(matrix) > 1 else [],
                "raw": matrix,
                "n_rows": len(matrix),
                "n_cols": max((len(r) for r in matrix), default=0),
            }
        )
        texts.append(name + "\n" + "\n".join("\t".join(r) for r in matrix[:80]))
    return {
        "title": path.stem,
        "published_at": None,
        "text": "\n\n".join(texts),
        "blocks": [{"ordinal": 0, "block_type": "xlsx", "heading": path.stem, "text": "\n\n".join(texts), "locator": "xlsx-body"}],
        "tables": tables,
        "doc_type": "xlsx",
    }


def parse_file(path: Path) -> dict[str, Any]:
    suf = path.suffix.lower()
    if suf in {".html", ".htm", ".shtml"}:
        return extract_html(path)
    if suf == ".pdf":
        return extract_pdf(path)
    if suf == ".docx":
        return extract_docx(path)
    if suf == ".doc":
        return extract_doc(path)
    if suf in {".xlsx", ".xlsm"}:
        return extract_xlsx(path)
    if suf == ".xls":
        return extract_xls(path)
    if suf in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        return extract_image(path, content_hash=path.stem)
    if suf == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace")
        return {
            "title": path.stem,
            "published_at": None,
            "text": text,
            "blocks": [{"ordinal": 0, "block_type": "text", "heading": path.stem, "text": text, "locator": "text"}],
            "tables": [],
            "doc_type": "txt",
        }
    return {
        "title": path.stem,
        "published_at": None,
        "text": "",
        "blocks": [],
        "tables": [],
        "doc_type": suf.lstrip(".") or "bin",
        "unparsed": True,
    }


def store_parsed(conn, url: str, raw_path: str, parsed: dict[str, Any], manifest: dict[str, Any]) -> int:
    title = parsed.get("title") or url
    if title and re.fullmatch(r"[0-9a-f-]{32,}", title):
        first = (parsed.get("text") or "").strip().splitlines()
        title = next((ln.strip() for ln in first if len(ln.strip()) >= 6), title)
    source = conn.execute("SELECT id, org, topics, notes FROM sources WHERE url=?", (url,)).fetchone()
    source_id = source["id"] if source else None
    org = source["org"] if source else None
    topics = source["topics"] if source else None
    parent_hint = (source["notes"] if source else None) or (topics.split(",")[0] if topics else "")
    generic = title in {"南京邮电大学", "南京邮电大学本科", "南京邮电大学  编", "南京邮电大学编", "附件", "正文", "通知"}
    parent_has_theme = parent_hint and any(
        k in parent_hint for k in ("就业质量", "培养方案", "招生章程", "教学质量", "学生手册", "保研", "推免", "工作方案")
    )
    title_has_theme = any(
        k in (title or "") for k in ("就业质量", "培养方案", "招生章程", "教学质量", "学生手册", "保研", "推免", "工作方案")
    )
    if parent_hint and (
        not title
        or title == url
        or generic
        or re.fullmatch(r"[0-9a-f-]{32,}", title)
        or len(title) < 6
        or (parent_has_theme and not title_has_theme)
    ):
        title = parent_hint[:80]
    if title == url or not title:
        from urllib.parse import unquote
        title = unquote(url.rsplit("/", 1)[-1])[:80]
    year = primary_year(title, parsed.get("text") or "")
    academic = extract_academic_year(f"{title} {parsed.get('text','')[:800]}")
    cur = conn.execute(
        """
        INSERT INTO documents(source_id, url, title, doc_type, org, topics, published_at,
            fetched_at, content_hash, raw_path, parse_status, quality_flags)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(url) DO UPDATE SET
            title=excluded.title,
            doc_type=excluded.doc_type,
            published_at=COALESCE(excluded.published_at, documents.published_at),
            fetched_at=excluded.fetched_at,
            content_hash=excluded.content_hash,
            raw_path=excluded.raw_path,
            parse_status=excluded.parse_status
        """,
        (
            source_id,
            url,
            title,
            parsed.get("doc_type"),
            org,
            topics,
            parsed.get("published_at"),
            manifest.get("fetched_at"),
            manifest.get("content_hash"),
            raw_path,
            "parsed" if not parsed.get("unparsed") else "unparsed",
            None,
        ),
    )
    doc_id = conn.execute("SELECT id FROM documents WHERE url=?", (url,)).fetchone()["id"]
    conn.execute(
        """
        INSERT OR IGNORE INTO document_versions(document_id, content_hash, fetched_at, raw_path, http_status)
        VALUES(?,?,?,?,?)
        """,
        (doc_id, manifest.get("content_hash"), manifest.get("fetched_at"), raw_path, manifest.get("http_status")),
    )
    version = conn.execute(
        "SELECT id FROM document_versions WHERE document_id=? AND content_hash=?",
        (doc_id, manifest.get("content_hash")),
    ).fetchone()
    version_id = version["id"] if version else None

    # replace old blocks/tables for this document
    old_tables = [r["id"] for r in conn.execute("SELECT id FROM tables WHERE document_id=?", (doc_id,))]
    for tid in old_tables:
        conn.execute("DELETE FROM table_cells WHERE table_id=?", (tid,))
        conn.execute("DELETE FROM table_rows WHERE table_id=?", (tid,))
    conn.execute("DELETE FROM tables WHERE document_id=?", (doc_id,))
    conn.execute("DELETE FROM blocks WHERE document_id=?", (doc_id,))
    # retrieval units for this doc
    old_ru = [r["id"] for r in conn.execute("SELECT id FROM retrieval_units WHERE document_id=?", (doc_id,))]
    for rid in old_ru:
        conn.execute("DELETE FROM retrieval_fts WHERE rowid=?", (rid,))
    conn.execute("DELETE FROM retrieval_units WHERE document_id=?", (doc_id,))

    for b in parsed.get("blocks") or []:
        tok = count_tokens(b.get("text") or "")
        conn.execute(
            """INSERT INTO blocks(document_id, version_id, ordinal, block_type, heading, text, locator, pdf_page, token_count)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (
                doc_id,
                version_id,
                b.get("ordinal"),
                b.get("block_type"),
                b.get("heading"),
                b.get("text"),
                b.get("locator"),
                b.get("pdf_page"),
                tok,
            ),
        )
        for chunk_i, chunk in enumerate(split_chunks(b.get("text") or "")):
            insert_retrieval_unit(
                conn,
                {
                    "document_id": doc_id,
                    "unit_type": "block",
                    "title": f"{title} / {b.get('heading') or ''}".strip(" /"),
                    "text": chunk,
                    "year": year,
                    "academic_year": academic,
                    "locator": f"{b.get('locator')}#c{chunk_i}",
                    "token_count": count_tokens(chunk),
                    "parent_locator": b.get("locator"),
                },
            )

    for tb in parsed.get("tables") or []:
        cur = conn.execute(
            """INSERT INTO tables(document_id, version_id, title, caption, locator, raw_json, header_json,
                n_rows, n_cols, pdf_page) VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (
                doc_id,
                version_id,
                tb.get("title"),
                tb.get("title"),
                tb.get("locator"),
                json.dumps(tb.get("raw") or [], ensure_ascii=False),
                json.dumps(tb.get("headers") or [], ensure_ascii=False),
                tb.get("n_rows"),
                tb.get("n_cols"),
                tb.get("pdf_page"),
            ),
        )
        table_id = int(cur.lastrowid)
        headers = tb.get("headers") or []
        raw = tb.get("raw") or []
        for ri, row in enumerate(raw):
            role = "header" if ri == tb.get("header_row", 0) else "body"
            joined = " | ".join(row)
            if any(x in joined for x in ("合计", "总计", "小计")):
                role = "total" if "小计" not in joined else "subtotal"
            conn.execute(
                "INSERT INTO table_rows(table_id, row_index, role, cells_json, text) VALUES(?,?,?,?,?)",
                (table_id, ri, role, json.dumps(row, ensure_ascii=False), joined),
            )
            for ci, val in enumerate(row):
                num = None
                m = re.search(r"-?\d+(?:\.\d+)?", val.replace(",", ""))
                if m and re.fullmatch(r"-?\d+(?:\.\d+)?%?", val.replace(",", "").replace(" ", "")):
                    try:
                        num = float(m.group())
                    except ValueError:
                        num = None
                header_path = headers[ci] if ci < len(headers) else f"col{ci}"
                conn.execute(
                    """INSERT INTO table_cells(table_id, row_index, col_index, header_path, value, numeric)
                       VALUES(?,?,?,?,?,?)""",
                    (table_id, ri, ci, header_path, val, num),
                )
        # retrieval: whole table + each body row with headers
        header_line = " | ".join(headers)
        whole = f"{title} {tb.get('title') or ''}\n{header_line}\n" + "\n".join(
            " | ".join(r) for r in (tb.get("rows") or [])
        )
        insert_retrieval_unit(
            conn,
            {
                "document_id": doc_id,
                "unit_type": "table",
                "title": f"{title} / {tb.get('title') or '表格'}",
                "text": whole[:6000],
                "year": year,
                "academic_year": academic,
                "locator": tb.get("locator"),
                "token_count": count_tokens(whole[:6000]),
            },
        )
        for ri, row in enumerate(tb.get("rows") or []):
            pairs = []
            for ci, val in enumerate(row):
                h = headers[ci] if ci < len(headers) else f"列{ci+1}"
                pairs.append(f"{h}={val}")
            row_text = f"{title} {tb.get('title') or ''} {year or ''}；" + "；".join(pairs)
            insert_retrieval_unit(
                conn,
                {
                    "document_id": doc_id,
                    "unit_type": "table_row",
                    "title": f"{title} 行{ri+1}",
                    "text": row_text,
                    "year": year,
                    "academic_year": academic,
                    "locator": f"{tb.get('locator')}#r{ri}",
                    "token_count": count_tokens(row_text),
                    "parent_locator": tb.get("locator"),
                },
            )
    return int(doc_id)


def parse_all(limit: int | None = None, redo_hashes: set[str] | None = None) -> dict:
    conn = init_db()
    rows = conn.execute(
        """
        SELECT url, content_hash, raw_path, fetched_at, http_status, ext, is_login_page, is_empty
        FROM crawl_manifest
        WHERE raw_path IS NOT NULL AND is_login_page=0 AND is_empty=0
          AND http_status BETWEEN 200 AND 399
        ORDER BY fetched_at DESC
        """
    ).fetchall()
    seen_hash: set[str] = set()
    already = {
        r["content_hash"]
        for r in conn.execute(
            """
            SELECT content_hash FROM documents
            WHERE parse_status IN ('parsed', 'discarded') AND content_hash IS NOT NULL
            """
        )
    }
    parsed = 0
    skipped = 0
    failed = 0
    for r in rows:
        if limit and parsed >= limit:
            break
        if r["content_hash"] in seen_hash or (r["content_hash"] in already and r["content_hash"] not in (redo_hashes or set())):
            skipped += 1
            continue
        seen_hash.add(r["content_hash"])
        path = Path(r["raw_path"])
        if not path.exists():
            failed += 1
            continue
        try:
            parsed_obj = parse_file(path)
            store_parsed(conn, r["url"], r["raw_path"], parsed_obj, dict(r))
            if parsed_obj.get("ocr") and (parsed_obj.get("ocr_conf") is not None and parsed_obj["ocr_conf"] < 55):
                add_quality(conn, "ocr_low_conf", f"{r['url']} conf={parsed_obj.get('ocr_conf')}")
            parsed += 1
            if parsed % 20 == 0:
                conn.commit()
        except Exception as e:
            failed += 1
            add_quality(conn, "parse_error", f"{r['url']}: {e}")
    conn.commit()
    conn.close()
    return {"parsed": parsed, "skipped_dup": skipped, "failed": failed}
